#!/usr/bin/env python3
"""
generate_excel_report.py
Professional Excel report generator for check_multi 2.0
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import json
import sys
from pathlib import Path

NAVY     = "1B365D"
SLATE    = "2F5496"
LIGHT    = "D6DCE4"
GREEN    = "C6EFCE"
GREEN_F  = "006100"
AMBER    = "FFEB9C"
AMBER_F  = "9C5700"
RED      = "FFC7CE"
RED_F    = "9C0006"
WHITE    = "FFFFFF"

thin = Border(
    left=Side(style='thin', color='B0B0B0'),
    right=Side(style='thin', color='B0B0B0'),
    top=Side(style='thin', color='B0B0B0'),
    bottom=Side(style='thin', color='B0B0B0')
)

def style_header(cell, dark=True):
    cell.font = Font(name='Calibri', bold=True, color=WHITE if dark else NAVY, size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY if dark else LIGHT)
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin

def auto_width(ws, min_width=10, max_width=45):
    for col in ws.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except Exception:
                pass
        adjusted = max(min_width, min(max_length + 2, max_width))
        ws.column_dimensions[column].width = adjusted

def build_report(summary: dict, output_path: Path):
    wb = Workbook()

    # Sheet 1 – Executive Dashboard
    ws = wb.active
    ws.title = "Executive Summary"

    ws.merge_cells('B2:G2')
    ws['B2'] = "DEVICE AUDIT REPORT"
    ws['B2'].font = Font(name='Calibri', bold=True, size=20, color=NAVY)
    ws['B2'].alignment = Alignment(horizontal='left', vertical='center')

    ws.merge_cells('B3:G3')
    ws['B3'] = f"{summary.get('check', 'unknown').upper()}  •  {summary.get('environment', 'lab').upper()}"
    ws['B3'].font = Font(name='Calibri', size=12, color=SLATE)

    meta = [
        ("Run ID", summary.get("run_id", "")),
        ("Timestamp", summary.get("timestamp", datetime.now().isoformat(timespec='seconds'))),
        ("Operator", summary.get("user", "")),
        ("Inventory", summary.get("inventory", "")),
        ("Tool", f"check_multi {summary.get('version', '2.0')}"),
    ]

    row = 5
    for label, value in meta:
        ws[f'B{row}'] = label
        ws[f'B{row}'].font = Font(bold=True, color=SLATE)
        ws[f'C{row}'] = value
        row += 1

    row = 11
    ws[f'B{row}'] = "RESULTS AT A GLANCE"
    ws[f'B{row}'].font = Font(bold=True, size=14, color=NAVY)

    kpis = [
        ("Total Devices", summary.get("total", 0), SLATE),
        ("Successful", summary.get("success", 0), GREEN_F),
        ("Failed", summary.get("failed", 0), RED_F),
        ("Success Rate", f"{summary.get('success_rate', 0):.1f}%", NAVY),
    ]

    col = 2
    for title, value, color in kpis:
        cell_title = ws.cell(row=row+1, column=col, value=title)
        cell_title.font = Font(size=9, color="666666")
        cell_value = ws.cell(row=row+2, column=col, value=value)
        cell_value.font = Font(bold=True, size=18, color=color)
        col += 2

    # Sheet 2 – Detailed Results
    ws2 = wb.create_sheet("Detailed Results")

    headers = ["Device", "Status", "Duration (s)", "Message / Notes"]
    for col_idx, header in enumerate(headers, 1):
        cell = ws2.cell(row=1, column=col_idx, value=header)
        style_header(cell)

    devices = summary.get("devices", [])
    for i, dev in enumerate(devices, 2):
        ws2.cell(row=i, column=1, value=dev.get("name", ""))
        status_cell = ws2.cell(row=i, column=2, value=dev.get("status", ""))
        ws2.cell(row=i, column=3, value=dev.get("duration", ""))
        ws2.cell(row=i, column=4, value=dev.get("message", ""))

        if dev.get("status") == "SUCCESS":
            status_cell.fill = PatternFill("solid", fgColor=GREEN)
            status_cell.font = Font(color=GREEN_F, bold=True)
        elif dev.get("status") == "FAILED":
            status_cell.fill = PatternFill("solid", fgColor=RED)
            status_cell.font = Font(color=RED_F, bold=True)
        else:
            status_cell.fill = PatternFill("solid", fgColor=AMBER)
            status_cell.font = Font(color=AMBER_F, bold=True)

        for c in range(1, 5):
            ws2.cell(row=i, column=c).border = thin
            ws2.cell(row=i, column=c).alignment = Alignment(vertical="center")

    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = f"A1:D{len(devices)+1}"

    auto_width(ws)
    auto_width(ws2)

    for sheet in [ws, ws2]:
        sheet.page_setup.orientation = 'landscape'
        sheet.page_setup.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = '1:1'

    wb.save(output_path)
    print(f"Excel report written → {output_path}")

if __name__ == "__main__":
    if len(sys.argv) < 3:
        print("Usage: generate_excel_report.py <summary.json> <output.xlsx>")
        sys.exit(1)
    with open(sys.argv[1]) as f:
        data = json.load(f)
    build_report(data, Path(sys.argv[2]))
